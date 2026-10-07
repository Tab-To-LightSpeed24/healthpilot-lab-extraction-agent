"""Backend support for the side-by-side document viewer and the live
processing feed: original-file serving, DOCX ingestion, content-type
correction, cheap list queries, and the progress log."""
import io
from unittest.mock import patch

import pytest
from docx import Document as DocxDocument
from sqlalchemy import event

from app.core.config import settings
from app.models.document import Document
from app.services import pipeline
from app.services.pdf_utils import DOCX_CONTENT_TYPE
from tests.test_api import client, process_pending, _make_pdf_bytes  # noqa: F401


@pytest.fixture(autouse=True)
def llm_off(monkeypatch):
    monkeypatch.setattr(settings, "llm_enabled", False)


def _docx_bytes() -> bytes:
    d = DocxDocument()
    d.add_paragraph("Acme Reference Laboratory")
    t = d.add_table(rows=3, cols=4)
    for r, row in enumerate([("Test", "Result", "Units", "Range"),
                             ("WBC", "6.8", "10*3/uL", "4.0-11.0"),
                             ("Hemoglobin", "13.9", "g/dL", "13.0-17.0")]):
        for c, text in enumerate(row):
            t.cell(r, c).text = text
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _process_without_outer_rollback(client, doc_id):
    """process_document's outer exception handler calls db.rollback(). On the
    shared test connection that would also roll back the test's own outer
    transaction (including the uploaded row) -- something a production session
    doesn't have -- so this stand-in makes that one call a no-op."""
    from sqlalchemy.orm import Session
    with patch.object(Session, "rollback", lambda self: None):
        process_pending(client, doc_id)


def _upload(client, name, data, content_type):
    resp = client.post("/reports", files={"file": (name, data, content_type)})
    assert resp.status_code == 201, resp.text
    return resp.json()


# ---- original file endpoint --------------------------------------------------

@pytest.mark.parametrize("name,data,ctype", [
    ("r.pdf", None, "application/pdf"),
    ("scan.png", b"\x89PNG\r\n\x1a\nfake", "image/png"),
    ("photo.jpg", b"\xff\xd8\xff\xe0fake", "image/jpeg"),
    ("r.txt", b"WBC 6.8 10*3/uL 4-11", "text/plain"),
    ("r.docx", None, DOCX_CONTENT_TYPE),
])
def test_original_file_is_served_back_byte_for_byte(client, name, data, ctype):
    data = data if data is not None else (_make_pdf_bytes("x") if name.endswith(".pdf") else _docx_bytes())
    doc = _upload(client, name, data, ctype)
    resp = client.get(f"/reports/{doc['id']}/file")
    assert resp.status_code == 200
    assert resp.content == data
    assert resp.headers["content-type"].startswith(ctype)
    assert resp.headers["content-disposition"].startswith("inline;")
    assert name in resp.headers["content-disposition"]


def test_file_endpoint_404_for_unknown_document(client):
    assert client.get("/reports/nope/file").status_code == 404


def test_non_ascii_filename_survives_the_download_header(client):
    doc = _upload(client, "rapport-été.txt", b"WBC 6.8 10*3/uL", "text/plain")
    resp = client.get(f"/reports/{doc['id']}/file")
    assert resp.status_code == 200 and "UTF-8''" in resp.headers["content-disposition"]


# ---- content-type correction --------------------------------------------------

@pytest.mark.parametrize("name,expected", [
    ("report.docx", DOCX_CONTENT_TYPE), ("report.PDF", "application/pdf"),
    ("scan.PNG", "image/png"), ("p.jpeg", "image/jpeg"), ("n.txt", "text/plain"),
])
def test_generic_octet_stream_is_resolved_from_the_extension(client, name, expected):
    body = _docx_bytes() if name.lower().endswith("docx") else b"data"
    doc = _upload(client, name, body, "application/octet-stream")
    assert doc["content_type"] == expected


def test_unknown_extension_with_generic_type_is_still_rejected(client):
    resp = client.post("/reports", files={"file": ("x.exe", b"MZ", "application/octet-stream")})
    assert resp.status_code == 415


# ---- list query stays cheap ----------------------------------------------------

def test_listing_reports_never_loads_the_stored_file_blob(client):
    _upload(client, "r.txt", b"WBC 6.8 10*3/uL 4-11", "text/plain")
    statements = []

    def capture(conn, cursor, statement, *a):
        statements.append(statement)

    engine = client.test_session_factory.kw["bind"].engine
    event.listen(engine, "before_cursor_execute", capture)
    try:
        assert client.get("/reports").status_code == 200
        _ = client.get(f"/reports/{client.get('/reports').json()[0]['id']}").json()
    finally:
        event.remove(engine, "before_cursor_execute", capture)
    selects = [s for s in statements if s.lstrip().upper().startswith("SELECT") and "FROM documents" in s]
    assert selects and all("raw_content" not in s for s in selects)


# ---- DOCX ingestion ------------------------------------------------------------

def test_docx_tables_are_extracted_without_any_llm(client):
    doc = _upload(client, "r.docx", _docx_bytes(), DOCX_CONTENT_TYPE)
    process_pending(client, doc["id"])
    detail = client.get(f"/reports/{doc['id']}").json()
    rows = {o["original_test_name"]: o for o in detail["observations"]}
    assert detail["status"] == "complete" and detail["num_pages"] == 1
    assert rows["WBC"]["value"] == "6.8" and rows["WBC"]["unit"] == "10*3/uL"
    assert rows["Hemoglobin"]["reference_range"] == "13.0-17.0"


def test_corrupt_docx_fails_with_a_clear_message(client):
    doc = _upload(client, "bad.docx", b"this is not a zip", DOCX_CONTENT_TYPE)
    _process_without_outer_rollback(client, doc["id"])
    detail = client.get(f"/reports/{doc['id']}").json()
    assert detail["status"] == "failed" and "could not read the .docx" in detail["error_message"]


def test_docx_is_sent_to_the_llm_as_text_when_it_is_enabled(client, monkeypatch):
    monkeypatch.setattr(settings, "llm_enabled", True)
    seen = {}

    def fake(image_png, text_layer):
        seen["image"], seen["text"] = image_png, text_layer
        from app.schemas.extraction import PageExtractionResult
        return PageExtractionResult(tests=[])

    doc = _upload(client, "r.docx", _docx_bytes(), DOCX_CONTENT_TYPE)
    with patch("app.services.pipeline.llm_client.extract_page", fake):
        process_pending(client, doc["id"])
    assert seen["image"] == b"" and "Hemoglobin" in seen["text"] and "13.9" in seen["text"]


# ---- live progress feed -----------------------------------------------------------

def test_a_queued_document_has_an_empty_feed(client):
    doc = _upload(client, "r.txt", b"WBC 6.8 10*3/uL 4-11", "text/plain")
    detail = client.get(f"/reports/{doc['id']}").json()
    assert detail["progress"] == [] and detail["pages_done"] == 0 and detail["current_step"] is None


def test_processing_writes_an_ordered_human_readable_feed(client):
    doc = _upload(client, "r.txt", b"WBC          6.8   10*3/uL  4.0-11.0\nHemoglobin   13.9  g/dL     13.0-17.0\n", "text/plain")
    process_pending(client, doc["id"])
    detail = client.get(f"/reports/{doc['id']}").json()
    msgs = [e["msg"] for e in detail["progress"]]

    expected_in_order = ["Picked up from the queue", "Opened r.txt", "AI extraction is turned off",
                         "Page 1/1: reading the text locally", "found 2 test row(s)",
                         "matching 2 row(s) to LOINC", "2 observation(s) extracted"]
    cursor = 0
    for needle in expected_in_order:
        idx = next((i for i in range(cursor, len(msgs)) if needle in msgs[i]), None)
        assert idx is not None, f"{needle!r} missing/out of order in {msgs}"
        cursor = idx + 1
    assert detail["pages_done"] == detail["num_pages"] == 1
    assert detail["current_step"] == msgs[-1]
    assert all(e["level"] in ("info", "warn", "error") and e["t"] for e in detail["progress"])
    assert client.get("/reports").json()[0]["current_step"] == msgs[-1]


def test_an_llm_failure_is_narrated_in_the_feed(client, monkeypatch):
    monkeypatch.setattr(settings, "llm_enabled", True)
    doc = _upload(client, "r.txt", b"WBC          6.8   10*3/uL  4.0-11.0\n", "text/plain")
    with patch("app.services.pipeline.llm_client.extract_page", side_effect=RuntimeError("402 insufficient credits")):
        process_pending(client, doc["id"])
    feed = client.get(f"/reports/{doc['id']}").json()["progress"]
    warn = [e for e in feed if e["level"] == "warn"]
    assert any("to the AI model in parallel" in e["msg"] for e in feed)
    assert any("AI unavailable" in e["msg"] and "insufficient credits" in e["msg"] for e in warn)


def test_a_failed_document_ends_its_feed_with_an_error(client):
    doc = _upload(client, "bad.docx", b"nope", DOCX_CONTENT_TYPE)
    _process_without_outer_rollback(client, doc["id"])
    feed = client.get(f"/reports/{doc['id']}").json()["progress"]
    assert feed[-1]["level"] == "error" and feed[-1]["msg"].startswith("Failed:")


def test_the_feed_is_capped_so_it_cannot_grow_without_bound(client):
    doc_row = _upload(client, "r.txt", b"x", "text/plain")
    session = client.test_session_factory()
    try:
        doc = session.query(Document).filter(Document.id == doc_row["id"]).one()
        for i in range(pipeline.MAX_PROGRESS_ENTRIES + 50):
            pipeline._progress(session, doc, f"step {i}")
        assert len(doc.progress) == pipeline.MAX_PROGRESS_ENTRIES
        assert doc.progress[-1]["msg"] == f"step {pipeline.MAX_PROGRESS_ENTRIES + 49}"
    finally:
        session.close()


def test_cancellation_is_recorded_in_the_feed(client):
    doc = _upload(client, "r.txt", b"WBC          6.8   10*3/uL  4.0-11.0\n", "text/plain")

    from app.services import pipeline as pl
    real = pl._is_cancel_requested
    with patch.object(pl, "_is_cancel_requested", lambda db, did: True):
        process_pending(client, doc["id"])
    detail = client.get(f"/reports/{doc['id']}").json()
    assert detail["status"] == "cancelled"
    assert detail["progress"][-1]["msg"].startswith("Cancelled")
